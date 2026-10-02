"""Apart gün sonu raporu Excel dosyalarını üretir.

Kullanım: python3 olustur.py
Çıktı:  Apart_Gunluk_Rapor_SABLON.xlsx  (boş, kullanıma hazır)
        Apart_Gunluk_Rapor_ORNEK_30.09.2026.xlsx  (orijinal dosyadaki veriyle dolu)
"""
from datetime import date, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

N = 1005            # Konaklama son satır (6..1005 = 1000 kayıt)
P = 505             # Ödemeler son satır (6..505 = 500 kayıt)
FIRST = 6
ODA_ILK, ODA_SON = 3, 32
MAX_KISI = 6
PAY_ROWS = 20

DATE_FMT = "dd.mm.yyyy"
MONEY_FMT = "#,##0.00"
thin = Side(style="thin", color="999999")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD = PatternFill("solid", fgColor="1F3864")
INPUT = PatternFill("solid", fgColor="FFF2CC")
CALC = PatternFill("solid", fgColor="E7E6E6")
SUB = PatternFill("solid", fgColor="D9E1F2")
WHITE_B = Font(bold=True, color="FFFFFF")
RED = PatternFill("solid", bgColor="FFC7CE", fgColor="FFC7CE")
GREEN = PatternFill("solid", bgColor="C6EFCE", fgColor="C6EFCE")
ORANGE = PatternFill("solid", bgColor="FCE4D6", fgColor="FCE4D6")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def name(wb, n, ref):
    wb.defined_names[n] = DefinedName(n, attr_text=ref)


def hdr(ws, row, labels, col=1):
    for i, t in enumerate(labels):
        c = ws.cell(row, col + i, t)
        c.fill, c.font, c.alignment, c.border = HEAD, WHITE_B, CENTER, BOX


def dv(ws, rng, **kw):
    v = DataValidation(allow_blank=True, showErrorMessage=True, errorStyle="stop", **kw)
    ws.add_data_validation(v)
    v.add(rng)
    return v


def build(path, rapor_tarihi, kirli, konaklama, odemeler):
    wb = Workbook()
    ws = wb.active
    ws.title = "Rapor"
    wk = wb.create_sheet("Konaklama")
    wp = wb.create_sheet("Ödemeler")
    wa = wb.create_sheet("Ayarlar")
    wh = wb.create_sheet("Hesap")
    wh.sheet_state = "hidden"

    # ---------------- Ayarlar ----------------
    wa["A1"], wa["A1"].font = "AYARLAR (yalnızca yönetici değiştirir)", Font(bold=True, size=13)
    hdr(wa, 3, ["Oda No", "", "Açıklama / Kaynak listesi", "", "Ödeme türü listesi", "", "Genel"])
    for i, o in enumerate(range(ODA_ILK, ODA_SON + 1)):
        wa.cell(4 + i, 1, o)
    for i, k in enumerate(["ÖĞRENCİ", "MÜNFERİT", "ODAMAX", "JOLLYTUR", "COMP-MARATON"]):
        wa.cell(4 + i, 3, k)
    for i, t in enumerate(["NAKİT", "KK", "HAVALE/EFT"]):
        wa.cell(4 + i, 5, t)
    wa["G4"], wa["H4"] = "Odada azami kişi", MAX_KISI
    wa["G5"] = "Yeni kaynak/ödeme türü eklemek için ilgili listenin altındaki boş hücreye yazın."
    for r in range(4, 4 + 30):
        for c in (1, 3, 5):
            wa.cell(r, c).fill = INPUT
    wa["H4"].fill = INPUT
    for col, w in zip("ACEGH", (10, 28, 22, 18, 10)):
        wa.column_dimensions[col].width = w
    name(wb, "OdaListesi", f"Ayarlar!$A$4:$A${3 + ODA_SON - ODA_ILK + 1}")
    name(wb, "KaynakListesi", "OFFSET(Ayarlar!$C$4,0,0,COUNTA(Ayarlar!$C$4:$C$33),1)")
    name(wb, "TurListesi", "OFFSET(Ayarlar!$E$4,0,0,COUNTA(Ayarlar!$E$4:$E$13),1)")
    name(wb, "KaynakAralik", "Ayarlar!$C$4:$C$33")
    name(wb, "TurAralik", "Ayarlar!$E$4:$E$13")
    name(wb, "MaxKisi", "Ayarlar!$H$4")
    name(wb, "RaporTarihi", "Rapor!$A$4")

    # ---------------- Konaklama ----------------
    wk["A1"] = "KONAKLAMA KAYITLARI"
    wk["A1"].font = Font(bold=True, size=14)
    wk["A2"] = ("Her konaklama için BİR satır girin (odadan çıkan misafir de silinmez, kayıtta kalır). "
                "Sarı hücreler girilir, gri hücreler otomatiktir. Kontrol sütunu 'OK' olmadan rapor onaylanmaz.")
    wk["A3"] = ("Giriş Durumu: Bekleniyor = bugün gelecek ama henüz gelmedi · Geldi = otele giriş yaptı · İptal = gelmedi/iptal.")
    wk["A2"].alignment = wk["A3"].alignment = Alignment(wrap_text=False)
    hdr(wk, 5, ["Oda No", "Misafir Adı", "Kişi Sayısı", "Giriş Tarihi", "Çıkış Tarihi", "Giriş Durumu",
                "Toplam Tutar", "Açıklama / Kaynak", "Süre (Gün)", "Günlük Oda Fiyatı", "Kontrol", "Gün Sonu Anahtar"])
    for r in range(FIRST, N + 1):
        for c in range(1, 9):
            cell = wk.cell(r, c)
            cell.fill, cell.border = INPUT, BOX
            cell.protection = Protection(locked=False)
        wk.cell(r, 4).number_format = wk.cell(r, 5).number_format = DATE_FMT
        wk.cell(r, 7).number_format = MONEY_FMT
        wk.cell(r, 9, f'=IF(AND(ISNUMBER(D{r}),ISNUMBER(E{r})),E{r}-D{r},"")')
        wk.cell(r, 10, f'=IF(AND(ISNUMBER(I{r}),ISNUMBER(G{r})),IF(I{r}>0,G{r}/I{r},""),"")').number_format = MONEY_FMT
        wk.cell(r, 11, (
            f'=IF(COUNTA(A{r}:H{r})=0,"",'
            f'IF(OR(A{r}="",B{r}="",C{r}="",D{r}="",E{r}="",F{r}="",G{r}="",H{r}=""),"HATA: eksik bilgi",'
            f'IF(COUNTIF(OdaListesi,A{r})=0,"HATA: geçersiz oda no",'
            f'IF(OR(NOT(ISNUMBER(D{r})),NOT(ISNUMBER(E{r}))),"HATA: tarih geçersiz",'
            f'IF(E{r}<=D{r},"HATA: çıkış girişten sonra olmalı",'
            f'IF(OR(NOT(ISNUMBER(C{r})),C{r}<1,C{r}>MaxKisi,C{r}<>INT(C{r})),"HATA: kişi sayısı geçersiz",'
            f'IF(OR(NOT(ISNUMBER(G{r})),G{r}<0),"HATA: tutar geçersiz",'
            f'IF(B{r}<>TRIM(B{r}),"HATA: isimde baştaki/sondaki boşluk",'
            f'IF(COUNTIF(KaynakAralik,H{r})=0,"HATA: kaynak listede yok",'
            f'IF(AND(F{r}<>"Bekleniyor",F{r}<>"Geldi",F{r}<>"İptal"),"HATA: giriş durumu geçersiz",'
            f'IF(AND(F{r}<>"İptal",COUNTIFS($A${FIRST}:$A${N},A{r},$D${FIRST}:$D${N},"<"&E{r},'
            f'$E${FIRST}:$E${N},">"&D{r},$F${FIRST}:$F${N},"<>İptal")>1),"HATA: oda bu tarihlerde çakışıyor",'
            f'IF(AND(F{r}="Bekleniyor",D{r}<RaporTarihi),"HATA: giriş günü geçti, Geldi veya İptal yapın",'
            f'"OK"))))))))))))'))
        wk.cell(r, 12, (f'=IF(AND(ISNUMBER(A{r}),F{r}<>"İptal",ISNUMBER(D{r}),ISNUMBER(E{r}),'
                        f'D{r}<=RaporTarihi,E{r}>RaporTarihi),A{r},"")'))
        for c in range(9, 13):
            wk.cell(r, c).fill = CALC
            wk.cell(r, c).border = BOX
    for col, w in zip("ABCDEFGHIJKL", (9, 32, 8, 13, 13, 14, 14, 20, 10, 14, 38, 8)):
        wk.column_dimensions[col].width = w
    wk.column_dimensions["L"].hidden = True
    wk.freeze_panes = "A6"
    wk.auto_filter.ref = f"A5:K{N}"
    rng = lambda col: f"{col}{FIRST}:{col}{N}"
    dv(wk, rng("A"), type="list", formula1="=OdaListesi", errorTitle="Geçersiz oda",
       error=f"Oda no {ODA_ILK}-{ODA_SON} arasında listeden seçilmelidir.")
    dv(wk, rng("B"), type="custom", formula1=f"=AND(LEN(B{FIRST})>=3,B{FIRST}=TRIM(B{FIRST}))",
       errorTitle="Geçersiz isim", error="İsim en az 3 harf olmalı; başında/sonunda boşluk olmamalı.")
    dv(wk, rng("C"), type="whole", operator="between", formula1="1", formula2="MaxKisi",
       errorTitle="Geçersiz kişi sayısı", error=f"Kişi sayısı 1 ile {MAX_KISI} arasında tam sayı olmalı.")
    dv(wk, rng("D"), type="date", operator="between", formula1="43831", formula2="73050",
       errorTitle="Geçersiz tarih", error="Giriş tarihi gerçek bir tarih olmalı (gg.aa.yyyy). Metin yazılamaz.")
    dv(wk, rng("E"), type="custom",
       formula1=(f"=AND(ISNUMBER(E{FIRST}),E{FIRST}>D{FIRST},COUNTIFS($A${FIRST}:$A${N},A{FIRST},"
                 f"$D${FIRST}:$D${N},\"<\"&E{FIRST},$E${FIRST}:$E${N},\">\"&D{FIRST},$F${FIRST}:$F${N},\"<>İptal\")<=1)"),
       errorTitle="Geçersiz çıkış tarihi",
       error="Çıkış tarihi girişten SONRA olmalı ve oda bu tarihlerde başka bir misafirde olmamalı. "
             "Önce oda no ve giriş tarihini girin.")
    dv(wk, rng("F"), type="list", formula1='"Bekleniyor,Geldi,İptal"', errorTitle="Geçersiz durum",
       error="Bekleniyor, Geldi veya İptal seçin.")
    dv(wk, rng("G"), type="decimal", operator="greaterThanOrEqual", formula1="0",
       errorTitle="Geçersiz tutar", error="Tutar 0 veya daha büyük bir sayı olmalı (comp için 0).")
    dv(wk, rng("H"), type="list", formula1="=KaynakListesi", errorTitle="Listede yok",
       error="Listeden seçin. Yeni kaynak için yöneticiye (Ayarlar sayfası) başvurun.")
    wk.conditional_formatting.add(f"A{FIRST}:K{N}", FormulaRule(formula=[f'LEFT($K{FIRST},4)="HATA"'], fill=RED))
    wk.conditional_formatting.add(f"A{FIRST}:K{N}", FormulaRule(formula=[f'$F{FIRST}="İptal"'], font=Font(strike=True, color="808080")))
    wk.conditional_formatting.add(f"F{FIRST}:F{N}", FormulaRule(formula=[f'$F{FIRST}="Bekleniyor"'], fill=ORANGE))
    wk.conditional_formatting.add(f"K{FIRST}:K{N}", FormulaRule(formula=[f'$K{FIRST}="OK"'], fill=GREEN))

    # ---------------- Ödemeler ----------------
    wp["A1"] = "ÖDEME KAYITLARI"
    wp["A1"].font = Font(bold=True, size=14)
    wp["A2"] = "Alınan her ödeme için bir satır. Tarih = ödemenin alındığı gün. Odada o tarihte misafir yoksa sistem kabul etmez."
    hdr(wp, 5, ["Tarih", "Oda No", "Ödemeyi Yapan", "Tutar", "Ödeme Türü", "Açıklama (isteğe bağlı)", "Kontrol", "Sıra"])
    for r in range(FIRST, P + 1):
        for c in range(1, 7):
            cell = wp.cell(r, c)
            cell.fill, cell.border = INPUT, BOX
            cell.protection = Protection(locked=False)
        wp.cell(r, 1).number_format = DATE_FMT
        wp.cell(r, 4).number_format = MONEY_FMT
        wp.cell(r, 7, (
            f'=IF(COUNTA(A{r}:F{r})=0,"",'
            f'IF(OR(A{r}="",B{r}="",C{r}="",D{r}="",E{r}=""),"HATA: eksik bilgi",'
            f'IF(OR(NOT(ISNUMBER(A{r})),NOT(ISNUMBER(D{r})),D{r}<=0),"HATA: tarih/tutar geçersiz",'
            f'IF(C{r}<>TRIM(C{r}),"HATA: isimde baştaki/sondaki boşluk",'
            f'IF(COUNTIF(TurAralik,E{r})=0,"HATA: ödeme türü geçersiz",'
            f'IF(COUNTIFS(kOda,B{r},kGiris,"<="&A{r},kCikis,">="&A{r},kDurum,"<>İptal")=0,'
            f'"HATA: bu odada o tarihte misafir yok","OK"))))))'))
        wp.cell(r, 8, f'=IF(AND(ISNUMBER(A{r}),A{r}=RaporTarihi),COUNTIFS($A${FIRST}:A{r},RaporTarihi),"")')
        for c in (7, 8):
            wp.cell(r, c).fill = CALC
            wp.cell(r, c).border = BOX
    for col, w in zip("ABCDEFGH", (13, 9, 32, 14, 14, 30, 38, 6)):
        wp.column_dimensions[col].width = w
    wp.column_dimensions["H"].hidden = True
    wp.freeze_panes = "A6"
    wp.auto_filter.ref = f"A5:G{P}"
    rp = lambda col: f"{col}{FIRST}:{col}{P}"
    dv(wp, rp("A"), type="date", operator="between", formula1="43831", formula2="73050",
       errorTitle="Geçersiz tarih", error="Gerçek bir tarih girin (gg.aa.yyyy). Metin yazılamaz.")
    dv(wp, rp("B"), type="list", formula1="=OdaListesi", errorTitle="Geçersiz oda", error="Listeden oda no seçin.")
    dv(wp, rp("C"), type="custom", formula1=f"=AND(LEN(C{FIRST})>=3,C{FIRST}=TRIM(C{FIRST}))",
       errorTitle="Geçersiz isim", error="İsim en az 3 harf olmalı; başında/sonunda boşluk olmamalı.")
    dv(wp, rp("D"), type="decimal", operator="greaterThan", formula1="0", errorTitle="Geçersiz tutar",
       error="Tutar 0'dan büyük bir sayı olmalı.")
    dv(wp, rp("E"), type="list", formula1="=TurListesi", errorTitle="Geçersiz tür", error="Listeden seçin.")
    wp.conditional_formatting.add(f"A{FIRST}:G{P}", FormulaRule(formula=[f'LEFT($G{FIRST},4)="HATA"'], fill=RED))
    wp.conditional_formatting.add(f"G{FIRST}:G{P}", FormulaRule(formula=[f'$G{FIRST}="OK"'], fill=GREEN))

    # ---------------- Hesap (gizli) ----------------
    hdr(wh, 1, ["Oda", "GünBaşı", "Gelen", "Gelecek", "Giden", "GünSonu", "Boş",
                "tGB", "tGel", "tGlc", "tGid", "tGS", "tBoş", "Sıra"])
    for i in range(ODA_SON - ODA_ILK + 1):
        r = 2 + i
        wh.cell(r, 1, f"=Ayarlar!A{4 + i}")
        wh.cell(r, 2, f'=IF(COUNTIFS(kOda,A{r},kGiris,"<"&RaporTarihi,kCikis,">="&RaporTarihi,kDurum,"Geldi")>0,1,0)')
        wh.cell(r, 3, f'=IF(COUNTIFS(kOda,A{r},kGiris,RaporTarihi,kDurum,"Geldi")>0,1,0)')
        wh.cell(r, 4, f'=IF(COUNTIFS(kOda,A{r},kGiris,RaporTarihi,kDurum,"Bekleniyor")>0,1,0)')
        wh.cell(r, 5, f'=IF(COUNTIFS(kOda,A{r},kCikis,RaporTarihi,kDurum,"Geldi")>0,1,0)')
        wh.cell(r, 6, f'=IF(COUNTIFS(kOda,A{r},kGiris,"<="&RaporTarihi,kCikis,">"&RaporTarihi,kDurum,"<>İptal")>0,1,0)')
        wh.cell(r, 7, f"=1-F{r}")
        for j, src in enumerate("BCDEFG"):
            wh.cell(r, 8 + j, f'=IF({src}{r}=1,TEXT($A{r},"00"),"")')
        wh.cell(r, 14, f'=IF(F{r}=1,COUNTIF(F$2:F{r},1),"")')
    last = 1 + ODA_SON - ODA_ILK + 1
    for nm, col in zip(["hGB", "hGel", "hGlc", "hGid", "hGS", "hBos"], "BCDEFG"):
        name(wb, nm, f"Hesap!${col}$2:${col}${last}")
    name(wb, "hSira", f"Hesap!$N$2:$N${last}")
    name(wb, "hOda", f"Hesap!$A$2:$A${last}")

    for nm, col in [("kOda", "A"), ("kAd", "B"), ("kKisi", "C"), ("kGiris", "D"), ("kCikis", "E"),
                    ("kDurum", "F"), ("kTutar", "G"), ("kKaynak", "H"), ("kSure", "I"), ("kGunluk", "J"),
                    ("kKontrol", "K"), ("kAnahtar", "L")]:
        name(wb, nm, f"Konaklama!${col}${FIRST}:${col}${N}")
    for nm, col in [("pTarih", "A"), ("pOda", "B"), ("pYapan", "C"), ("pTutar", "D"), ("pTur", "E"),
                    ("pAcik", "F"), ("pKontrol", "G"), ("pSira", "H")]:
        name(wb, nm, f"'Ödemeler'!${col}${FIRST}:${col}${P}")

    # ---------------- Rapor ----------------
    for col, w in zip("ABCDEFGHI", (13, 30, 14, 14, 14, 16, 18, 17, 24)):
        ws.column_dimensions[col].width = w
    ws.column_dimensions["K"].hidden = True
    ws.merge_cells("A1:I1")
    ws["A1"] = "APART GÜN SONU RAPORU"
    ws["A1"].font, ws["A1"].alignment = Font(bold=True, size=16), CENTER
    ws.merge_cells("A2:I2")
    ws["A2"] = (
        '=IF(COUNTIF(kKontrol,"HATA*")+COUNTIF(pKontrol,"HATA*")>0,'
        '"⚠ DİKKAT: "&(COUNTIF(kKontrol,"HATA*")+COUNTIF(pKontrol,"HATA*"))&" kayıtta HATA var — Konaklama/Ödemeler sayfalarındaki kırmızı satırları düzeltin",'
        'IF(OR(F5<>SUM(hGS),F7<>SUMIFS(kKisi,kGiris,"<="&RaporTarihi,kCikis,">"&RaporTarihi,kDurum,"<>İptal")),'
        '"⚠ DİKKAT: gün sonu oda/kişi sayısı tutmuyor — kayıtları kontrol edin",'
        f'IF(COUNT(pSira)>{PAY_ROWS},"⚠ DİKKAT: {PAY_ROWS}\'den fazla ödeme var, rapora sığmıyor",'
        'IF(I6="","⚠ Kirli boş oda sayısını giriniz (yoksa 0 yazın)","✔ Kontroller temiz — rapor hazır"))))')
    ws["A2"].font, ws["A2"].alignment = Font(bold=True, size=12), CENTER
    ws.row_dimensions[2].height = 24
    ws.conditional_formatting.add("A2:I2", FormulaRule(formula=['LEFT($A$2,1)="⚠"'], fill=RED, font=Font(bold=True, color="9C0006")))
    ws.conditional_formatting.add("A2:I2", FormulaRule(formula=['LEFT($A$2,1)="✔"'], fill=GREEN, font=Font(bold=True, color="006100")))

    ws.merge_cells("A4:A5")
    ws["A4"] = rapor_tarihi
    ws["A4"].number_format, ws["A4"].fill = DATE_FMT, INPUT
    ws["A4"].font, ws["A4"].alignment = Font(bold=True, size=12), CENTER
    ws["A4"].protection = Protection(locked=False)
    for c, t in zip("BCDEFG", ["Gün Başı", "Gelen", "Gelecek", "Giden", "Gün Sonu", "Boş Odalar"]):
        ws[f"{c}4"] = t
    hdr(ws, 4, ["Gün Başı", "Gelen", "Gelecek", "Giden", "Gün Sonu", "Boş Odalar"], col=2)
    ws.merge_cells("H4:I4")
    ws["H4"] = "BOŞ ODA DURUMU"
    ws["H4"].fill, ws["H4"].font, ws["H4"].alignment = HEAD, WHITE_B, CENTER
    for c, nm in zip("BCDE", ["hGB", "hGel", "hGlc", "hGid"]):
        ws[f"{c}5"] = f"=SUM({nm})"
    ws["F5"] = "=B5-E5+C5+D5"
    ws["G5"] = "=COUNT(OdaListesi)-F5"
    ws.merge_cells("H5:I5")
    ws["H5"] = '=IF(I6="","",(G5-I6)&" ODA BOŞ VE TEMİZ // "&I6&" ODA BOŞ VE KİRLİ")'
    ws["A6"], ws["A7"] = "Oda No", "Kişi Sayısı"
    ws["B6"] = '=_xlfn.TEXTJOIN("-",TRUE,Hesap!H2:H%d)' % last
    ws["C6"] = '=_xlfn.TEXTJOIN("-",TRUE,Hesap!I2:I%d)' % last
    ws["D6"] = '=_xlfn.TEXTJOIN("-",TRUE,Hesap!J2:J%d)' % last
    ws["E6"] = '=_xlfn.TEXTJOIN("-",TRUE,Hesap!K2:K%d)' % last
    ws["F6"] = '=_xlfn.TEXTJOIN("-",TRUE,Hesap!L2:L%d)' % last
    ws["G6"] = '=_xlfn.TEXTJOIN("-",TRUE,Hesap!M2:M%d)' % last
    ws["H6"] = "Kirli boş oda sayısı ►"
    ws["I6"] = kirli
    ws["I6"].fill, ws["I6"].protection = INPUT, Protection(locked=False)
    ws["B7"] = '=SUMIFS(kKisi,kGiris,"<"&RaporTarihi,kCikis,">="&RaporTarihi,kDurum,"Geldi")'
    ws["C7"] = '=SUMIFS(kKisi,kGiris,RaporTarihi,kDurum,"Geldi")'
    ws["D7"] = '=SUMIFS(kKisi,kGiris,RaporTarihi,kDurum,"Bekleniyor")'
    ws["E7"] = '=SUMIFS(kKisi,kCikis,RaporTarihi,kDurum,"Geldi")'
    ws["F7"] = "=B7-E7+C7+D7"
    ws.merge_cells("H7:I7")
    ws["H7"] = '=G5&" ODA BOŞ"'
    ws.row_dimensions[6].height = 78
    for r in (4, 5, 6, 7):
        for c in "ABCDEFGHI":
            cell = ws[f"{c}{r}"]
            cell.border = BOX
            if r > 4 and c != "A" or r == 4 and c == "A":
                cell.alignment = CENTER
            if r in (6, 7) and c == "A":
                cell.fill, cell.font, cell.alignment = SUB, Font(bold=True), CENTER
    ws["F5"].font = ws["F7"].font = Font(bold=True)
    dv(ws, "A4", type="date", operator="between", formula1="43831", formula2="73050",
       errorTitle="Geçersiz tarih", error="Gerçek bir tarih girin (gg.aa.yyyy).")
    dv(ws, "I6", type="whole", operator="between", formula1="0", formula2="$G$5",
       errorTitle="Geçersiz sayı", error="Kirli oda sayısı 0 ile boş oda sayısı arasında tam sayı olmalı.")

    # Gün sonu tablosu
    TH, T0 = 9, 10
    hdr(ws, TH, ["Oda - Gün Sonu", "İsim - Gün Sonu", "Kişi Sayısı", "Giriş Tarihi", "Çıkış Tarihi",
                 "Süre (Gün)", "Günlük Oda Fiyatı", "Toplam Tutar", "Açıklama"])
    n_rows = ODA_SON - ODA_ILK + 1
    T1 = T0 + n_rows - 1
    for i in range(n_rows):
        r = T0 + i
        ws[f"A{r}"] = f'=IFERROR(INDEX(hOda,MATCH({i + 1},hSira,0)),"")'
        ws[f"K{r}"] = f'=IF(A{r}="","",MATCH(A{r},kAnahtar,0))'
        for col, nm in zip("BCDEFGH", ["kAd", "kKisi", "kGiris", "kCikis", "kSure", "kGunluk", "kTutar"]):
            ws[f"{col}{r}"] = f'=IF($K{r}="","",INDEX({nm},$K{r}))'
        ws[f"I{r}"] = f'=IF($K{r}="","",INDEX(kKaynak,$K{r})&IF(INDEX(kDurum,$K{r})="Bekleniyor"," (GELECEK)",""))'
        ws[f"D{r}"].number_format = ws[f"E{r}"].number_format = DATE_FMT
        ws[f"G{r}"].number_format = ws[f"H{r}"].number_format = MONEY_FMT
        for col in "ABCDEFGHI":
            ws[f"{col}{r}"].border = BOX
            ws[f"{col}{r}"].alignment = CENTER if col != "B" else LEFT
    ws.conditional_formatting.add(f"A{T0}:I{T1}", FormulaRule(formula=[f'ISNUMBER(SEARCH("GELECEK",$I{T0}))'], fill=ORANGE))

    r = T1 + 2
    for c, t in (("A", "Toplam Oda"), ("C", "Toplam Kişi"), ("G", "Günlük Toplam"), ("H", "Günlük Ort")):
        ws[f"{c}{r}"] = t
        ws[f"{c}{r}"].fill, ws[f"{c}{r}"].font, ws[f"{c}{r}"].alignment, ws[f"{c}{r}"].border = SUB, Font(bold=True), CENTER, BOX
    ws[f"A{r + 1}"] = f"=COUNT(A{T0}:A{T1})"
    ws[f"C{r + 1}"] = f"=SUM(C{T0}:C{T1})"
    ws[f"G{r + 1}"] = f"=SUM(G{T0}:G{T1})"
    ws[f"H{r + 1}"] = f"=IFERROR(G{r + 1}/A{r + 1},0)"
    for c in "AGH":
        pass
    ws[f"G{r + 1}"].number_format = ws[f"H{r + 1}"].number_format = MONEY_FMT
    for c in "ACGH":
        ws[f"{c}{r + 1}"].border, ws[f"{c}{r + 1}"].alignment, ws[f"{c}{r + 1}"].font = BOX, CENTER, Font(bold=True)

    # Ödemeler bölümü
    q = r + 3
    hdr(ws, q, ["ÖDEME - ODA", "ÖDEMEYİ YAPAN", "TUTAR", "ÖDEME TÜRÜ", "AÇIKLAMA"])
    ws.merge_cells(f"E{q}:I{q}")
    for i in range(PAY_ROWS):
        rr = q + 1 + i
        ws[f"K{rr}"] = f'=IFERROR(MATCH({i + 1},pSira,0),"")'
        ws[f"A{rr}"] = f'=IF($K{rr}="","",INDEX(pOda,$K{rr}))'
        ws[f"B{rr}"] = f'=IF($K{rr}="","",INDEX(pYapan,$K{rr}))'
        ws[f"C{rr}"] = f'=IF($K{rr}="","",INDEX(pTutar,$K{rr}))'
        ws[f"D{rr}"] = f'=IF($K{rr}="","",INDEX(pTur,$K{rr}))'
        ws[f"E{rr}"] = f'=IF($K{rr}="","",INDEX(pAcik,$K{rr})&"")'
        ws.merge_cells(f"E{rr}:I{rr}")
        ws[f"C{rr}"].number_format = MONEY_FMT
        for col in "ABCDEFGHI":
            ws[f"{col}{rr}"].border = BOX
            ws[f"{col}{rr}"].alignment = CENTER if col != "B" else LEFT
    s = q + PAY_ROWS + 2
    ws[f"A{s}"], ws[f"C{s}"] = "TOPLAM TAHSİLAT", f'=SUMIFS(pTutar,pTarih,RaporTarihi)'
    for k in range(4):
        ws[f"A{s + 1 + k}"] = f'=IF(Ayarlar!E{4 + k}="","",Ayarlar!E{4 + k})'
        ws[f"C{s + 1 + k}"] = f'=IF(A{s + 1 + k}="","",SUMIFS(pTutar,pTarih,RaporTarihi,pTur,A{s + 1 + k}))'
    for rr in range(s, s + 5):
        ws[f"C{rr}"].number_format = MONEY_FMT
        for col in "AB":
            ws[f"{col}{rr}"].border = BOX
        ws[f"C{rr}"].border = BOX
        ws[f"A{rr}"].font = ws[f"C{rr}"].font = Font(bold=(rr == s))
        ws[f"A{rr}"].fill = SUB if rr == s else PatternFill()
    ws.merge_cells(f"A{s}:B{s}")
    for k in range(1, 5):
        ws.merge_cells(f"A{s + k}:B{s + k}")

    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f"A1:I{s + 4}"
    ws.sheet_view.showGridLines = False

    for sheet in (ws, wh):
        sheet.protection.sheet = True
    for sheet in (wk, wp):
        sheet.protection.sheet = True
        sheet.protection.autoFilter = False
        sheet.protection.formatColumns = False
        sheet.protection.formatRows = False
        sheet.protection.sort = False
    wb.calculation.fullCalcOnLoad = True

    # Örnek veri
    for i, row in enumerate(konaklama):
        for c, v in enumerate(row, 1):
            wk.cell(FIRST + i, c, v)
    for i, row in enumerate(odemeler):
        for c, v in enumerate(row, 1):
            wp.cell(FIRST + i, c, v)
    wb.save(path)


D = date
def stay(oda, ad, kisi, giris, sure, tutar, kaynak, durum="Geldi"):
    return [oda, ad, kisi, giris, giris + timedelta(days=sure), durum, tutar, kaynak]

ORNEK_K = [
    stay(18, "MUHAMMET SEMİH KÖSE", 4, D(2026, 9, 13), 285, 498192, "ÖĞRENCİ"),
    stay(20, "MEHMET EFE ÇELEBİ", 4, D(2026, 9, 21), 267, 471522, "ÖĞRENCİ"),
    stay(22, "HAYAL YAVAŞ", 4, D(2026, 9, 16), 272, 480352, "ÖĞRENCİ"),
    stay(24, "BEREN TEZSEZENER", 2, D(2026, 9, 12), 276, 482298, "ÖĞRENCİ"),
    stay(26, "HİLAL KIVRAK", 2, D(2026, 9, 4), 284, 496450, "ÖĞRENCİ"),
    stay(29, "EDA NUR RUŞİTLER", 2, D(2026, 9, 13), 275, 480532, "ÖĞRENCİ"),
    stay(13, "ZAKARIA EL KHATABI", 3, D(2026, 9, 25), 6, 30600, "MÜNFERİT"),
    stay(16, "SALİH ARABAYAPAN", 5, D(2026, 9, 27), 7, 37999.99, "MÜNFERİT"),
    stay(31, "MEHMET EMIN DINC", 4, D(2026, 9, 26), 7, 37087, "ODAMAX"),
    stay(8, "KÜBRA TOPRAK KOLCU", 4, D(2026, 9, 30), 1, 4337, "JOLLYTUR"),
    stay(14, "BATI BAKKALOĞLU", 5, D(2026, 9, 30), 4, 20000, "MÜNFERİT"),
    stay(25, "MEHMET YENER", 4, D(2026, 9, 30), 4, 2000, "COMP-MARATON", "Bekleniyor"),
    stay(28, "ÇERKEZ BİLİR", 2, D(2026, 9, 30), 4, 2000, "COMP-MARATON", "Bekleniyor"),
    # Orijinal dosyada sadece oda no / kişi sayısı vardı (06-07-08 giden, toplam 10 kişi); isim ve tutar örnektir.
    stay(6, "ÖRNEK GİDEN MİSAFİR 06", 3, D(2026, 9, 27), 3, 0, "MÜNFERİT"),
    stay(7, "ÖRNEK GİDEN MİSAFİR 07", 3, D(2026, 9, 27), 3, 0, "MÜNFERİT"),
    stay(8, "ÖRNEK GİDEN MİSAFİR 08", 4, D(2026, 9, 27), 3, 0, "MÜNFERİT"),
]
ORNEK_P = [[D(2026, 9, 30), 14, "BATI BAKKALOĞLU", 10000, "KK", ""]]

if __name__ == "__main__":
    import os
    here = os.path.dirname(os.path.abspath(__file__))
    build(os.path.join(here, "Apart_Gunluk_Rapor_SABLON.xlsx"), D(2026, 10, 1), None, [], [])
    build(os.path.join(here, "Apart_Gunluk_Rapor_ORNEK_30.09.2026.xlsx"), D(2026, 9, 30), 1, ORNEK_K, ORNEK_P)
